#!/usr/bin/env python3
"""Build the pupil-import spreadsheet template.

The columns here are not invented — they are exactly what
apps/ems/src/lib/import.ts matches on, and the rules on the instructions
sheet are what apps/ems/src/lib/actions/import-actions.ts actually does with
each row. If the importer changes, change this too.

    python3 tools/make-import-template.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT_XLSX = ROOT / "docs" / "templates" / "pupil-import-template.xlsx"
OUT_CSV = ROOT / "docs" / "templates" / "pupil-import-template.csv"
EMS_PUBLIC = ROOT / "apps" / "ems" / "public" / "pupil-import-template.xlsx"

# The school's twelve classes, as seeded. normalizeClassName() is forgiving —
# "Grade 1", "GRADE1", "G1" and "1" all resolve — but the dropdown offers the
# canonical spelling so nobody has to rely on that.
CLASSES = ["Playgroup", "PP1", "PP2"] + [f"Grade {n}" for n in range(1, 10)]

# (header, width, required, help)
COLUMNS = [
    ("Admission No", 16, False,
     "Leave blank and the system assigns the next one (SSK-26xxxx). Fill it in only "
     "to keep a number the pupil already has."),
    ("UPI / Assessment No", 20, False,
     "The learner's NEMIS/KNEC number, if you have it. Used to spot a pupil who is "
     "already in the system."),
    ("Surname", 16, True,
     "Family name. Required."),
    ("First Name", 22, True,
     "Given name(s) — put middle names here too. Required."),
    ("Gender", 10, True,
     "Male or Female. REQUIRED: a row with this blank is skipped. M/F, Boy/Girl also work."),
    ("Class", 14, True,
     "Which class the pupil joins. Must already exist under Academics → Classes. Required."),
    ("Stream", 10, False,
     "A, B, C… Leave blank and the pupil goes into the class's first stream."),
    ("Boarding", 12, False,
     "Day or Boarder. Leave blank and the pupil is recorded as a day scholar."),
    ("Guardian Name", 24, False,
     "One parent or guardian. Creates the guardian record attached to the pupil."),
    ("Guardian Phone", 18, False,
     "The number the office would ring. Format it as text (07… keeps its leading zero)."),
]

EXAMPLE_ROWS = [
    ["", "", "Wekesa", "Amani John", "Male", "Grade 1", "A", "Day", "Mary Wekesa", "0712000001"],
    ["", "", "Nafula", "Blessing", "Female", "Grade 1", "A", "Day", "Peter Nafula", "0712000002"],
    ["SSK-260041", "", "Atieno", "Neema", "Female", "Playgroup", "", "", "James Atieno", "0712000014"],
    ["", "12345678", "Barasa", "Faith Nekesa", "Female", "Grade 7", "A", "Boarder", "Alice Barasa", "0712000021"],
]

INK = "1F2937"
BRAND = "1D6FA3"
RULE = "D8DCE3"
HEAD_FILL = PatternFill("solid", fgColor="1D6FA3")
REQ_FILL = PatternFill("solid", fgColor="FFF4E5")
EXAMPLE_FILL = PatternFill("solid", fgColor="F3F6F9")
thin = Side(style="thin", color=RULE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def build() -> None:
    wb = Workbook()

    # ------------------------------------------------------------- Pupils --
    ws = wb.active
    ws.title = "Pupils"
    ws.freeze_panes = "A2"

    for i, (header, width, required, _help) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i, value=header)
        c.font = Font(bold=True, color="FFFFFF", size=11)
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 24

    # Example rows, visibly marked so nobody imports them by accident.
    for r, row in enumerate(EXAMPLE_ROWS, start=2):
        for i, value in enumerate(row, start=1):
            c = ws.cell(row=r, column=i, value=value)
            c.fill = EXAMPLE_FILL
            c.font = Font(italic=True, color="6B7280")
            c.border = BOX
            if COLUMNS[i - 1][0] in ("Guardian Phone", "Admission No", "UPI / Assessment No"):
                c.number_format = "@"

    first_blank = 2 + len(EXAMPLE_ROWS)
    last = first_blank + 1200

    # Keep phone numbers and admission numbers as text: Excel otherwise eats
    # the leading zero off 0712000001 and the office gets 712000001.
    for col_name in ("Admission No", "UPI / Assessment No", "Guardian Phone"):
        idx = [c[0] for c in COLUMNS].index(col_name) + 1
        letter = get_column_letter(idx)
        for r in range(first_blank, last + 1):
            ws.cell(row=r, column=idx).number_format = "@"

    def add_validation(col_name: str, values: list[str], prompt: str) -> None:
        idx = [c[0] for c in COLUMNS].index(col_name) + 1
        letter = get_column_letter(idx)
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(values) + '"',
            allow_blank=True,
            showDropDown=False,
        )
        dv.promptTitle = col_name
        dv.prompt = prompt
        dv.errorTitle = "Pick one from the list"
        dv.error = prompt
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{last}")

    add_validation("Gender", ["Male", "Female"], "Male or Female. A row without this is skipped.")
    add_validation("Class", CLASSES, "Must be a class that exists under Academics → Classes.")
    add_validation("Boarding", ["Day", "Boarder"], "Blank counts as Day.")
    add_validation("Stream", ["A", "B", "C", "D"], "Blank uses the class's first stream.")

    # Shade the required columns all the way down, so a half-filled row is
    # obvious on screen before anyone uploads it.
    for i, (_h, _w, required, _help) in enumerate(COLUMNS, start=1):
        if not required:
            continue
        for r in range(first_blank, last + 1):
            ws.cell(row=r, column=i).fill = REQ_FILL

    # -------------------------------------------------- How to fill this in --
    doc = wb.create_sheet("How to fill this in")
    doc.column_dimensions["A"].width = 24
    doc.column_dimensions["B"].width = 12
    doc.column_dimensions["C"].width = 86

    def heading(text: str, row: int) -> int:
        c = doc.cell(row=row, column=1, value=text)
        c.font = Font(bold=True, size=13, color=BRAND)
        return row + 1

    def para(text: str, row: int) -> int:
        c = doc.cell(row=row, column=1, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        doc.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        doc.row_dimensions[row].height = 15 * (1 + len(text) // 110)
        return row + 1

    r = 1
    r = heading("Bringing the register in", r)
    r = para(
        "Type the pupils into the Pupils sheet, one per row. Delete the four grey example "
        "rows before you import — they are there to show the shape, not to be admitted.",
        r,
    )
    r += 1
    r = heading("What each column is for", r)
    for name, _w, required, help_text in COLUMNS:
        doc.cell(row=r, column=1, value=name).font = Font(bold=True)
        req = doc.cell(row=r, column=2, value="required" if required else "optional")
        req.font = Font(color="B45309" if required else "6B7280", bold=required)
        h = doc.cell(row=r, column=3, value=help_text)
        h.alignment = Alignment(wrap_text=True, vertical="top")
        doc.row_dimensions[r].height = 15 * (1 + len(help_text) // 90)
        r += 1

    r += 1
    r = heading("Importing it", r)
    for step in [
        "1.  File → Save As → CSV (Comma delimited). The importer reads CSV, not .xlsx.",
        "2.  In the system: Students → Import pupils from a spreadsheet.",
        "3.  Choose the CSV. You get a PREVIEW first — every row, and what will happen to it.",
        "4.  Read the preview. Anything that cannot be admitted says why, by row number.",
        "5.  Only then press the button to admit them.",
    ]:
        r = para(step, r)

    r += 1
    r = heading("Things worth knowing before you start", r)
    for note in [
        "An academic session must be active (Settings → Sessions & terms), or the import "
        "refuses — pupils have to be enrolled into something.",
        "A class must already exist under Academics → Classes. The import will not invent "
        "one; a row naming an unknown class is skipped and says so.",
        "Gender is required. A row without it is skipped. Everything else except the name "
        "and the class can be blank.",
        "Admission numbers are assigned automatically when you leave the column blank.",
        "Re-running the same file is safe. A pupil already in the system — matched on "
        "admission number or UPI — is skipped, not duplicated.",
        "Up to 1,000 rows per file. Split a bigger register and import it in parts.",
        "Names are tidied to Title Case. WEKESA AMANI becomes Wekesa Amani.",
        "If you have one 'Learner Name' column instead of Surname + First Name, that works "
        "too — it is read surname-first, the KNEC way.",
        "Keep phone numbers as text so 0712000001 does not lose its leading zero. This "
        "template already does that; watch for it if you paste from elsewhere.",
    ]:
        r = para("•  " + note, r)

    r += 1
    r = heading("What happens to each pupil", r)
    r = para(
        "They are admitted into the class and stream you named, enrolled into the active "
        "session, and given an admission number. A guardian name creates a guardian record "
        "attached to the pupil, with the phone number if you supplied one. Nothing else is "
        "assumed — fees, transport and boarding placement are set separately.",
        r,
    )

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    EMS_PUBLIC.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    wb.save(EMS_PUBLIC)

    # A CSV with the headers only, for anyone who would rather not open Excel.
    OUT_CSV.write_text(",".join(c[0] for c in COLUMNS) + "\n", encoding="utf-8")

    print(f"wrote {OUT_XLSX.relative_to(ROOT)}")
    print(f"wrote {EMS_PUBLIC.relative_to(ROOT)}")
    print(f"wrote {OUT_CSV.relative_to(ROOT)}")


if __name__ == "__main__":
    build()
